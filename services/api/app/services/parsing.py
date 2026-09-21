from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Any


@dataclass
class ParsedPage:
    page_number: int
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)
    image_bytes: bytes | None = None


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]
    metadata: dict[str, Any] = field(default_factory=dict)


def _table_payload(
    rows: list[list[Any]],
    *,
    table_index: int,
    max_rows: int = 1000,
) -> dict[str, Any]:
    normalized = [
        ["" if cell is None else str(cell).strip() for cell in row]
        for row in rows
    ]
    row_count = len(normalized)
    column_count = max((len(row) for row in normalized), default=0)
    return {
        "table_index": table_index,
        "rows": normalized[:max_rows],
        "row_count": row_count,
        "column_count": column_count,
        "truncated": row_count > max_rows,
    }


def _parse_pdf(data: bytes) -> ParsedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is required to parse PDFs") from exc

    reader = PdfReader(io.BytesIO(data))
    pages: list[ParsedPage] = []
    fitz_pdf = None
    try:
        import fitz

        fitz_pdf = fitz.open(stream=data, filetype="pdf")
    except Exception:
        fitz_pdf = None

    try:
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            image_bytes = None
            tables: list[dict[str, Any]] = []

            if fitz_pdf is not None:
                try:
                    fitz_page = fitz_pdf.load_page(i - 1)
                    finder = fitz_page.find_tables()
                    for table_index, table in enumerate(finder.tables, start=1):
                        extracted = table.extract()
                        if extracted:
                            tables.append(
                                _table_payload(
                                    extracted,
                                    table_index=table_index,
                                    max_rows=500,
                                )
                            )
                except Exception:
                    tables = []

            if len(text.strip()) < 20 and fitz_pdf is not None:
                try:
                    fitz_page = fitz_pdf.load_page(i - 1)
                    pix = fitz_page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    image_bytes = pix.tobytes("png")
                except Exception:
                    image_bytes = None

            pages.append(
                ParsedPage(
                    page_number=i,
                    text=text,
                    metadata={
                        "kind": "pdf_page",
                        "tables": tables,
                        "table_count": len(tables),
                    },
                    image_bytes=image_bytes,
                )
            )
    finally:
        if fitz_pdf is not None:
            fitz_pdf.close()

    return ParsedDocument(
        pages=pages,
        metadata={"title": reader.metadata.title if reader.metadata else None},
    )


def _parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    blocks: list[str] = []
    tables: list[dict[str, Any]] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table_index, table in enumerate(doc.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            blocks.append("\n".join(" | ".join(row) for row in rows))
            tables.append(_table_payload(rows, table_index=table_index))
    return ParsedDocument(
        [
            ParsedPage(
                1,
                "\n\n".join(blocks),
                {
                    "kind": "docx",
                    "tables": tables,
                    "table_count": len(tables),
                },
            )
        ]
    )


def _parse_pptx(data: bytes) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    pages: list[ParsedPage] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        tables: list[dict[str, Any]] = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                rows = [
                    [cell.text.strip() for cell in row.cells]
                    for row in shape.table.rows
                ]
                if rows:
                    tables.append(
                        _table_payload(rows, table_index=len(tables) + 1)
                    )
                    texts.append("\n".join(" | ".join(row) for row in rows))
            elif hasattr(shape, "text") and str(shape.text).strip():
                texts.append(str(shape.text).strip())
        pages.append(
            ParsedPage(
                i,
                "\n".join(texts),
                {
                    "kind": "slide",
                    "slide_number": i,
                    "tables": tables,
                    "table_count": len(tables),
                },
            )
        )
    return ParsedDocument(pages)


def _parse_csv(data: bytes) -> ParsedDocument:
    decoded = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(decoded)))
    rendered = "\n".join(" | ".join(cell for cell in row) for row in rows)
    table = _table_payload(rows, table_index=1, max_rows=5000) if rows else None
    return ParsedDocument(
        [
            ParsedPage(
                1,
                rendered,
                {
                    "kind": "csv",
                    "row_count": len(rows),
                    "tables": [table] if table else [],
                    "table_count": 1 if table else 0,
                },
            )
        ]
    )


def parse_document(data: bytes, mime_type: str) -> ParsedDocument:
    if mime_type == "application/pdf":
        return _parse_pdf(data)
    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _parse_docx(data)
    if mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        return _parse_pptx(data)
    if mime_type == "text/csv":
        return _parse_csv(data)
    if mime_type.startswith("text/"):
        return ParsedDocument(
            [ParsedPage(1, data.decode("utf-8-sig", errors="replace"), {"kind": "text"})]
        )
    if mime_type.startswith("image/"):
        return ParsedDocument(
            [ParsedPage(1, "", {"kind": "image"}, image_bytes=data)]
        )
    raise ValueError(f"Unsupported parser MIME type: {mime_type}")
